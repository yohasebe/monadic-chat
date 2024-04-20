#!/usr/bin/env python

import argparse
import json
import os
from docx import Document
import openpyxl
from pptx import Presentation

def extract_text_docx(doc_path):
    doc = Document(doc_path)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_xlsx(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path)
    text = []
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    text.append(str(cell.value))
    return "\n".join(text)

def extract_text_pptx(ppt_path):
    prs = Presentation(ppt_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

def export_as_json(file_path):
    file_type = os.path.splitext(file_path)[1].lower()
    if file_type == '.docx':
        text = extract_text_docx(file_path)
    elif file_type == '.xlsx':
        text = extract_text_xlsx(file_path)
    elif file_type == '.pptx':
        text = extract_text_pptx(file_path)
    else:
        raise ValueError("Unsupported file type")

    data = {'text': text}
    json_data = json.dumps(data, ensure_ascii=False, indent=4)
    print(json_data)

    base_filename = os.path.splitext(os.path.basename(file_path))[0]
    output_filename = f"{base_filename}.json"
    with open(output_filename, 'w', encoding='utf-8') as f:
        f.write(json_data)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Extract text from Office files and output as JSON.")
    parser.add_argument("file_path", help="Path to the Office file")
    args = parser.parse_args()

    if not os.path.exists(args.file_path):
        print(f"The specified file could not be found: {args.file_path}")
    else:
        export_as_json(args.file_path)
